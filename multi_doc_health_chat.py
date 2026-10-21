"""Loads ALL pdf/txt/docx from ./data, builds a FRESH index each run,
answers strictly from those files. pip install pypdf python-docx langchain
langchain-openai langchain-community chromadb python-dotenv"""

import os, glob, uuid
from dotenv import load_dotenv
from pypdf import PdfReader
from docx import Document as DocxDocument
from langchain_openai import OpenAIEmbeddings, ChatOpenAI
from langchain_community.vectorstores import Chroma
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.output_parsers import StrOutputParser
from langchain_core.runnables import RunnablePassthrough

load_dotenv()

EMB_BASE  = "https://ai-gateway.vercel.sh/v1"
EMB_KEY   = os.getenv("EMBEDDING_API_KEY")
EMB_MODEL = os.getenv("EMBEDDING_MODEL", "text-embedding-3-small")

LLM_BASE  = "https://api.novita.ai/openai"
LLM_KEY   = os.getenv("NOVITA_API_KEY")
LLM_MODEL = os.getenv("LLM_MODEL")   # set the model name you already use

DATA_DIR = "data"

def load_documents():
    docs = []
    for path in glob.glob(os.path.join(DATA_DIR, "**", "*.*"), recursive=True):
        ext = os.path.splitext(path)[1].lower()
        try:
            if ext == ".pdf":
                text = "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
            elif ext in (".txt", ".md"):
                text = open(path, encoding="utf-8", errors="ignore").read()
            elif ext == ".docx":
                text = "\n".join(p.text for p in DocxDocument(path).paragraphs)
            else:
                continue
            if text.strip():
                docs.append((text, os.path.basename(path)))
        except Exception as e:
            print(f"[WARN] could not read {path}: {e}")
    return docs

def build_index():
    raw = load_documents()
    print(f"[DEBUG] files loaded : {len(raw)} -> {[s for _, s in raw]}")
    assert raw, f"No readable documents found in ./{DATA_DIR}"
    splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
    chunks = [Document(page_content=c, metadata={"source": src})
              for text, src in raw for c in splitter.split_text(text)]
    print(f"[DEBUG] chunks created: {len(chunks)}")
    emb = OpenAIEmbeddings(api_key=EMB_KEY, base_url=EMB_BASE, model=EMB_MODEL)
    vs = Chroma.from_documents(chunks, emb,
                               collection_name=f"run-{uuid.uuid4().hex[:8]}")
    print(f"[DEBUG] vectors in DB: {vs._collection.count()}")
    return vs

def format_docs(docs):
    return "\n\n".join(f"[source: {d.metadata['source']}]\n{d.page_content}" for d in docs)

PROMPT = ChatPromptTemplate.from_template(
    """You are a helpful Health Assistant. Answer ONLY from the context of the user's documents.
If the answer is not there, reply exactly: "I don't have this information in your documents."
Add: "This is not medical advice."

Context:
{context}

Question: {question}""")

if __name__ == "__main__":
    vs = build_index()
    retriever = vs.as_retriever(search_kwargs={"k": 4})
    llm = ChatOpenAI(api_key=LLM_KEY, base_url=LLM_BASE, model=LLM_MODEL, temperature=0.1)
    chain = ({"context": retriever | format_docs, "question": RunnablePassthrough()}
             | PROMPT | llm | StrOutputParser())

    print("\nHealth Assistant ready. Type 'exit' to quit.")
    while True:
        q = input("\nYou: ").strip()
        if not q or q.lower() in {"exit", "quit"}:
            break
        hits = retriever.invoke(q)
        print(f"[sources used] {sorted({h.metadata['source'] for h in hits})}")
        print(f"\nAssistant: {chain.invoke(q)}")